import io
import os
import pytest
import docx
from app.parser import (
    extract_text_from_txt,
    extract_text_from_docx,
    extract_text_from_pdf,
    clean_extracted_text,
    audit_text_layer_integrity,
    audit_layout_linearization
)

SAMPLE_PDF_PATH = os.path.join(os.path.dirname(__file__), "..", "..", "sample_resume.pdf")

def test_extract_text_from_txt():
    content = b"Software Engineer with 5+ years of experience in Python, FastAPI, and React."
    text = extract_text_from_txt(content)
    assert "Software Engineer" in text
    assert "FastAPI" in text

def test_extract_text_from_txt_empty():
    content = b""
    with pytest.raises(ValueError, match="No text could be extracted"):
        extract_text_from_txt(content)

def test_extract_text_from_txt_utf8():
    content = "Senior Developer \u2014 Full Stack \u2022 Cloud Architecture".encode("utf-8")
    text = extract_text_from_txt(content)
    assert "Senior Developer" in text

def test_clean_extracted_text():
    dirty = "   Multiple   spaces \t and \n\n\n extra   lines. \x00 \u200B"
    cleaned = clean_extracted_text(dirty)
    assert "\x00" not in cleaned
    assert "\u200B" not in cleaned
    assert "Multiple spaces" in cleaned

def test_audit_text_layer_integrity_clean_text():
    clean_text = "Experienced Full Stack Engineer with expertise in Python, React, and Kubernetes."
    audit = audit_text_layer_integrity(clean_text)
    assert audit["text_layer_health_score"] == 100
    assert audit["is_searchable"] is True
    assert audit["pua_glyph_count"] == 0
    assert audit["replacement_char_count"] == 0
    assert len(audit["issues"]) == 0

def test_audit_text_layer_integrity_anomalous_text():
    corrupted_text = "Software Engineer \uE005 \uFFFD \u200B\u200B\u200B\u200B\u200B\u200B at ScaleFlow Systems"
    audit = audit_text_layer_integrity(corrupted_text)
    assert audit["pua_glyph_count"] == 1
    assert audit["replacement_char_count"] == 1
    assert audit["zero_width_char_count"] == 6
    assert audit["text_layer_health_score"] < 100
    assert len(audit["issues"]) >= 2

def test_audit_text_layer_integrity_empty():
    audit = audit_text_layer_integrity("")
    assert audit["text_layer_health_score"] == 0
    assert audit["is_searchable"] is False
    assert len(audit["issues"]) == 1

def test_extract_text_from_docx_with_paragraphs_and_tables():
    doc = docx.Document()
    doc.add_heading("Alex Rivera - Principal Engineer", 0)
    doc.add_paragraph("Spearheaded migration to microservices architecture improving latency by 45%.")
    
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Skills"
    table.cell(0, 1).text = "Python, Go, Docker"
    table.cell(1, 0).text = "Cloud"
    table.cell(1, 1).text = "AWS, GCP, Terraform"
    
    docx_io = io.BytesIO()
    doc.save(docx_io)
    docx_bytes = docx_io.getvalue()
    
    extracted = extract_text_from_docx(docx_bytes)
    assert "Alex Rivera" in extracted
    assert "Spearheaded migration" in extracted
    assert "Python, Go, Docker" in extracted

def test_extract_text_from_docx_empty_raises():
    empty_doc = docx.Document()
    docx_io = io.BytesIO()
    empty_doc.save(docx_io)
    empty_bytes = docx_io.getvalue()
    
    with pytest.raises(ValueError, match="No text could be extracted"):
        extract_text_from_docx(empty_bytes)

def test_extract_text_from_docx_max_chars_limit():
    doc = docx.Document()
    doc.add_paragraph("A" * 500)
    docx_io = io.BytesIO()
    doc.save(docx_io)
    docx_bytes = docx_io.getvalue()
    
    extracted = extract_text_from_docx(docx_bytes, max_chars=100)
    assert len(extracted) <= 100

def test_extract_text_from_real_pdf():
    if os.path.exists(SAMPLE_PDF_PATH):
        with open(SAMPLE_PDF_PATH, "rb") as f:
            pdf_bytes = f.read()
        extracted_text, page_count = extract_text_from_pdf(pdf_bytes)
        assert len(extracted_text) > 0
        assert page_count >= 1


def test_audit_layout_linearization_clean_single_column():
    clean_resume = (
        "Sarah Jenkins\n"
        "Lead Software Engineer\n"
        "sarah@example.com | (555) 234-5678\n"
        "Experience\n"
        "Spearheaded cloud migration to Kubernetes across 8 microservices.\n"
        "Automated CI/CD deployment pipelines using GitHub Actions.\n"
        "Education\n"
        "B.S. in Computer Science - University of California\n"
    )
    audit = audit_layout_linearization(clean_resume)
    assert audit["linearization_score"] == 100
    assert audit["risk_tier"] == "Safe Single-Column"
    assert audit["is_linear_safe"] is True
    assert audit["gutter_anomaly_lines"] == 0


def test_audit_layout_linearization_multi_column_gutter():
    multi_col_resume = (
        "SKILLS                      WORK EXPERIENCE\n"
        "Python, React, Docker       Senior Engineer at TechCorp\n"
        "PostgreSQL, Redis           Led distributed system migration\n"
        "AWS, Terraform, CI/CD       Reduced latency by 45%\n"
        "Kafka, RabbitMQ             Managed 5 direct report engineers\n"
    )
    audit = audit_layout_linearization(multi_col_resume)
    assert audit["gutter_anomaly_lines"] >= 4
    assert audit["linearization_score"] < 75
    assert "multi-column" in audit["issues"][0].lower()


def test_audit_layout_linearization_ascii_tables():
    table_resume = (
        "+-----------------------------------------------+\n"
        "| Project Name     | Tech Stack  | Outcome      |\n"
        "+-----------------------------------------------+\n"
        "| Stream Engine    | Python, Go  | 50k RPS      |\n"
        "+-----------------------------------------------+\n"
        "| Payment Gateway  | FastAPI     | $2M processed|\n"
        "+-----------------------------------------------+\n"
    )
    audit = audit_layout_linearization(table_resume)
    assert audit["divider_count"] >= 4
    assert audit["linearization_score"] < 100



def test_audit_layout_linearization_empty():
    assert audit_layout_linearization("")["linearization_score"] == 0
    assert audit_layout_linearization("   \n\t  ")["risk_tier"] == "Empty Document"

